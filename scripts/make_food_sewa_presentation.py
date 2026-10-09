from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


DESKTOP = Path(r"C:\Users\Alfa\OneDrive\Desktop")
IMAGE_DIR = DESKTOP / "food_sewa_doc_images"
OUTPUT = DESKTOP / "FOOD_SEWA_Presentation_from_doc.pptx"

WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

NAVY = RGBColor(9, 21, 43)
TEAL = RGBColor(15, 140, 128)
MINT = RGBColor(223, 248, 242)
LIGHT = RGBColor(247, 250, 252)
GRAY = RGBColor(96, 111, 126)
ORANGE = RGBColor(242, 144, 44)
GREEN = RGBColor(41, 151, 97)
BLUE = RGBColor(54, 103, 190)
RED = RGBColor(200, 61, 61)
WHITE = RGBColor(255, 255, 255)


def set_text(tf, text, size=18, bold=False, color=NAVY, align=None):
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    run = p.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.65))
    set_text(box.text_frame, title, 26, True, NAVY)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(0.98), Inches(1.25), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(1.95), Inches(0.82), Inches(10.6), Inches(0.32))
        set_text(sb.text_frame, subtitle, 11, False, GRAY)


def add_footer(slide, slide_no):
    box = slide.shapes.add_textbox(Inches(11.65), Inches(7.06), Inches(1.1), Inches(0.25))
    set_text(box.text_frame, f"{slide_no:02d}", 10, False, GRAY, PP_ALIGN.RIGHT)


def add_bullets(slide, items, x, y, w, h, font_size=17, color=NAVY):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(8)
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
    return box


def add_card(slide, x, y, w, h, title, body=None, fill=WHITE, accent=TEAL):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(215, 226, 231)
    shape.line.width = Pt(1)
    accent_bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, Inches(0.08), h)
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()
    tb = slide.shapes.add_textbox(x + Inches(0.22), y + Inches(0.15), w - Inches(0.35), Inches(0.32))
    set_text(tb.text_frame, title, 14, True, NAVY)
    if body:
        bb = slide.shapes.add_textbox(x + Inches(0.22), y + Inches(0.55), w - Inches(0.35), h - Inches(0.65))
        add_text = bb.text_frame
        add_text.word_wrap = True
        add_text.clear()
        p = add_text.paragraphs[0]
        p.text = body
        p.font.name = "Aptos"
        p.font.size = Pt(11.5)
        p.font.color.rgb = GRAY
        p.space_after = Pt(0)
    return shape


def add_image_fit(slide, image_name, x, y, w, h):
    path = IMAGE_DIR / image_name
    if not path.exists():
        return None
    return slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def add_node(slide, x, y, w, h, text, fill, font=13, text_color=WHITE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(shape.text_frame, text, font, True, text_color, PP_ALIGN.CENTER)
    return shape


def connect(slide, a, b, color=GRAY):
    ax = a.left + a.width
    ay = a.top + a.height // 2
    bx = b.left
    by = b.top + b.height // 2
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ax, ay, bx, by)
    line.line.color.rgb = color
    line.line.width = Pt(1.8)
    line.line.end_arrowhead = True
    return line


def add_table(slide, x, y, w, h, headers, rows, font_size=10.5):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w, h)
    table = table_shape.table
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL
        cell.text = header
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(font_size)
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(font_size)
                p.font.color.rgb = NAVY
                p.alignment = PP_ALIGN.LEFT
    return table_shape


def make_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = LIGHT
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


def main():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = MINT
    add_image_fit(slide, "image1.png", Inches(0.7), Inches(0.45), Inches(1.15), Inches(1.05))
    title = slide.shapes.add_textbox(Inches(0.85), Inches(2.0), Inches(7.6), Inches(1.2))
    set_text(title.text_frame, "FOOD SEWA", 48, True, NAVY)
    subtitle = slide.shapes.add_textbox(Inches(0.9), Inches(3.08), Inches(7.4), Inches(0.75))
    set_text(subtitle.text_frame, "Online Food Ordering and Delivery Management System", 20, False, TEAL)
    details = slide.shapes.add_textbox(Inches(0.92), Inches(4.08), Inches(6.2), Inches(1.0))
    set_text(details.text_frame, "Project Presentation\nAnil Prajapati | Rohit Dangol\nBCA, Samriddhi College", 15, False, NAVY)
    add_image_fit(slide, "image11.png", Inches(7.5), Inches(1.35), Inches(5.1), Inches(3.25))
    add_footer(slide, 1)

    # 2. Project Overview
    slide = make_slide(prs, "Project Overview", "A web-based platform connecting customers, restaurant operations, and delivery tracking")
    add_card(slide, Inches(0.75), Inches(1.55), Inches(3.7), Inches(1.45), "Customer Experience", "Register, browse food, add items to cart, place orders, pay online, and track delivery.")
    add_card(slide, Inches(4.85), Inches(1.55), Inches(3.7), Inches(1.45), "Admin Operations", "Manage menu items, categories, users, orders, payments, delivery status, and restaurant settings.", accent=ORANGE)
    add_card(slide, Inches(8.95), Inches(1.55), Inches(3.55), Inches(1.45), "Delivery Visibility", "Route-based status tracking with customer-facing order progress and estimated delivery information.", accent=GREEN)
    add_bullets(slide, [
        "Built with PHP, MySQL, HTML, CSS, JavaScript and Bootstrap.",
        "Designed for responsive use across desktop and mobile browsers.",
        "Focuses on secure authentication, order processing, payment integration, and delivery management.",
    ], Inches(1.05), Inches(3.65), Inches(11.1), Inches(1.5), 17)

    # 3. Problem Statement
    slide = make_slide(prs, "Problem Statement")
    add_bullets(slide, [
        "Traditional phone-based food ordering is slow, error-prone, and difficult to track.",
        "Small and medium restaurants often lack affordable digital order management.",
        "Customers need clearer menus, flexible payment options, and delivery visibility.",
        "Existing platforms can be expensive, less customizable, and limited outside core urban areas.",
    ], Inches(0.9), Inches(1.35), Inches(6.3), Inches(4.6), 19)
    a = add_node(slide, Inches(8.0), Inches(1.4), Inches(2.8), Inches(0.65), "Manual Ordering", RED)
    b = add_node(slide, Inches(8.0), Inches(2.45), Inches(2.8), Inches(0.65), "Order Errors", ORANGE)
    c = add_node(slide, Inches(8.0), Inches(3.5), Inches(2.8), Inches(0.65), "Delayed Updates", BLUE)
    d = add_node(slide, Inches(8.0), Inches(4.55), Inches(2.8), Inches(0.65), "Weak Analytics", GRAY)
    target = add_node(slide, Inches(10.9), Inches(2.75), Inches(1.6), Inches(1.15), "FOOD\nSEWA", TEAL, 17)
    for n in [a, b, c, d]:
        connect(slide, n, target)

    # 4. Objectives
    slide = make_slide(prs, "Objectives")
    rows = [
        ("01", "Real-time order tracking and delivery management"),
        ("02", "Digital menu management for restaurant operations"),
        ("03", "Online ordering with cart, address, and payment support"),
        ("04", "Role-based access for customer and admin workflows"),
        ("05", "Secure, maintainable PHP/MySQL implementation"),
    ]
    for i, (num, text) in enumerate(rows):
        y = Inches(1.28 + i * 1.02)
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(1.2), y, Inches(0.55), Inches(0.55))
        circle.fill.solid()
        circle.fill.fore_color.rgb = TEAL
        circle.line.fill.background()
        set_text(circle.text_frame, num, 12, True, WHITE, PP_ALIGN.CENTER)
        add_card(slide, Inches(1.95), y - Inches(0.05), Inches(9.7), Inches(0.68), text, None, WHITE, TEAL)

    # 5. Scope and Users
    slide = make_slide(prs, "Scope and System Users")
    roles = [
        ("Customer", "Registration, login, menu browsing, cart, checkout, order history, tracking"),
        ("Admin", "Dashboard, categories, foods, users, payments, orders, delivery management"),
        ("Delivery Flow", "Order status, tracking number, route progress, delivery completion"),
    ]
    for idx, (role, body) in enumerate(roles):
        add_card(slide, Inches(0.85 + idx * 4.1), Inches(1.35), Inches(3.55), Inches(1.45), role, body, WHITE, [TEAL, ORANGE, GREEN][idx])
    add_bullets(slide, [
        "Includes food ordering, payment gateway flow, address management, and route-based tracking.",
        "Does not cover native mobile apps, real rider GPS hardware, or full multi-restaurant marketplace deployment.",
        "GPS and payment behavior depends on third-party services and environment configuration.",
    ], Inches(1.0), Inches(3.75), Inches(11.2), Inches(1.65), 17)

    # 6. Methodology
    slide = make_slide(prs, "Development Methodology", "Agile-inspired incremental development")
    add_image_fit(slide, "image2.png", Inches(0.85), Inches(1.4), Inches(3.85), Inches(2.85))
    add_bullets(slide, [
        "Database-first: design schema and relationships for each module.",
        "Backend-second: implement PHP logic, validation, sessions, and CRUD operations.",
        "Frontend-third: build responsive Bootstrap interfaces and user workflows.",
        "Test each module before moving to the next feature increment.",
    ], Inches(5.25), Inches(1.45), Inches(6.9), Inches(3.35), 18)
    add_image_fit(slide, "image3.png", Inches(2.1), Inches(5.2), Inches(9.2), Inches(1.35))

    # 7. Architecture
    slide = make_slide(prs, "System Architecture")
    client1 = add_node(slide, Inches(0.9), Inches(2.0), Inches(2.2), Inches(0.75), "Customer\nWeb App", BLUE)
    client2 = add_node(slide, Inches(0.9), Inches(3.1), Inches(2.2), Inches(0.75), "Admin\nDashboard", BLUE)
    app = add_node(slide, Inches(4.1), Inches(1.55), Inches(3.05), Inches(0.72), "PHP Application Layer", TEAL, 15)
    modules = [
        add_node(slide, Inches(3.75), Inches(2.65), Inches(1.45), Inches(0.58), "Auth", GREEN, 11),
        add_node(slide, Inches(5.35), Inches(2.65), Inches(1.45), Inches(0.58), "Menu", GREEN, 11),
        add_node(slide, Inches(3.75), Inches(3.45), Inches(1.45), Inches(0.58), "Cart", GREEN, 11),
        add_node(slide, Inches(5.35), Inches(3.45), Inches(1.45), Inches(0.58), "Orders", GREEN, 11),
        add_node(slide, Inches(4.55), Inches(4.25), Inches(1.45), Inches(0.58), "Payment", GREEN, 11),
    ]
    db = add_node(slide, Inches(8.0), Inches(2.45), Inches(2.35), Inches(1.0), "MySQL\nDatabase", NAVY, 15)
    ext = add_node(slide, Inches(10.85), Inches(2.45), Inches(1.7), Inches(1.0), "External\nServices", ORANGE, 13)
    for c in [client1, client2]:
        connect(slide, c, app)
    connect(slide, app, db)
    connect(slide, db, ext, ORANGE)
    add_card(slide, Inches(8.05), Inches(4.2), Inches(4.45), Inches(1.35), "Integrations", "Stripe / eSewa / Khalti for payment, map services for location and route visibility.", WHITE, ORANGE)

    # 8. Data Model
    slide = make_slide(prs, "Data Model and Database Design")
    add_image_fit(slide, "image5.png", Inches(0.65), Inches(1.25), Inches(5.85), Inches(2.45))
    add_image_fit(slide, "image10.jpeg", Inches(7.0), Inches(1.1), Inches(2.55), Inches(4.85))
    add_bullets(slide, [
        "Users can have multiple addresses and place multiple orders.",
        "Orders contain many order items, preserving item name and price snapshots.",
        "Foods belong to categories; payments link to orders.",
        "Foreign keys and indexes support data integrity and efficient retrieval.",
    ], Inches(0.9), Inches(4.25), Inches(5.45), Inches(1.85), 15)
    add_card(slide, Inches(9.9), Inches(1.5), Inches(2.65), Inches(2.0), "Core Tables", "users, addresses, categories, foods, orders, order_items, payments, coupons", WHITE, TEAL)

    # 9. Process Flow
    slide = make_slide(prs, "Order Processing Flow")
    labels = ["Register/Login", "Browse Menu", "Add to Cart", "Checkout", "Payment", "Order Tracking"]
    nodes = []
    for idx, label in enumerate(labels):
        nodes.append(add_node(slide, Inches(0.65 + idx * 2.05), Inches(2.55), Inches(1.55), Inches(0.75), label, [BLUE, TEAL, GREEN, ORANGE, NAVY, TEAL][idx], 11))
    for left, right in zip(nodes, nodes[1:]):
        connect(slide, left, right)
    add_bullets(slide, [
        "Customer creates account and browses active menu items.",
        "Cart stores selected items and validates quantity against stock.",
        "Checkout validates saved address and selected payment gateway.",
        "Order receives a unique order code and can be tracked by customer/admin.",
    ], Inches(1.1), Inches(4.25), Inches(10.9), Inches(1.5), 16)

    # 10. DFD and Use Case
    slide = make_slide(prs, "Process and Use Case Diagrams")
    add_image_fit(slide, "image6.png", Inches(0.6), Inches(1.22), Inches(5.7), Inches(1.45))
    add_image_fit(slide, "image8.png", Inches(6.65), Inches(1.15), Inches(2.65), Inches(4.55))
    add_image_fit(slide, "image7.png", Inches(9.6), Inches(1.45), Inches(2.65), Inches(2.65))
    add_card(slide, Inches(0.85), Inches(3.2), Inches(5.15), Inches(2.2), "DFD Summary", "Customer, admin/restaurant, and delivery rider interact with the system through user management, menu management, order processing, payment processing, and delivery management.", WHITE, TEAL)

    # 11. Interface Design
    slide = make_slide(prs, "Interface Design")
    add_image_fit(slide, "image11.png", Inches(0.65), Inches(1.25), Inches(3.7), Inches(1.75))
    add_image_fit(slide, "image14.png", Inches(4.8), Inches(1.25), Inches(3.7), Inches(1.75))
    add_image_fit(slide, "image23.png", Inches(8.95), Inches(1.25), Inches(3.7), Inches(1.75))
    add_image_fit(slide, "image26.png", Inches(2.4), Inches(3.85), Inches(4.2), Inches(2.25))
    add_image_fit(slide, "image29.png", Inches(7.15), Inches(3.85), Inches(4.2), Inches(2.25))

    # 12. Algorithms
    slide = make_slide(prs, "Tracking and Algorithm Details")
    add_card(slide, Inches(0.9), Inches(1.35), Inches(3.75), Inches(1.45), "Haversine Distance", "Calculates real-world distance between restaurant and customer GPS coordinates.", WHITE, TEAL)
    add_card(slide, Inches(4.95), Inches(1.35), Inches(3.75), Inches(1.45), "Route Progress", "Uses order status and delivery timing to estimate progress and ETA.", WHITE, ORANGE)
    add_card(slide, Inches(9.0), Inches(1.35), Inches(3.25), Inches(1.45), "Late Coupon Logic", "Detects delayed deliveries and makes reward coupons claimable when eligible.", WHITE, GREEN)
    s = add_node(slide, Inches(1.25), Inches(4.0), Inches(1.8), Inches(0.68), "Restaurant GPS", BLUE)
    c = add_node(slide, Inches(4.0), Inches(4.0), Inches(1.8), Inches(0.68), "Customer GPS", BLUE)
    d = add_node(slide, Inches(6.75), Inches(4.0), Inches(1.8), Inches(0.68), "Distance + ETA", TEAL)
    t = add_node(slide, Inches(9.5), Inches(4.0), Inches(1.8), Inches(0.68), "Tracking UI", GREEN)
    connect(slide, s, c)
    connect(slide, c, d)
    connect(slide, d, t)

    # 13. Testing
    slide = make_slide(prs, "Testing Summary")
    add_table(slide, Inches(0.65), Inches(1.25), Inches(5.95), Inches(2.5),
              ["Unit", "Operation", "Result"],
              [
                  ("TC001", "Customer registration", "Pass"),
                  ("TC002", "Customer login", "Pass"),
                  ("TC003", "Add menu item", "Pass"),
                  ("TC004", "Place order", "Pass"),
              ], 10)
    add_table(slide, Inches(6.9), Inches(1.25), Inches(5.75), Inches(3.2),
              ["System", "Scenario", "Result"],
              [
                  ("ST001", "Customer order flow", "Pass"),
                  ("ST002", "Order status update", "Pass"),
                  ("ST003", "Delivery tracking", "Pass"),
                  ("ST004", "Cancellation/refund", "Pass"),
                  ("ST005", "Menu item visibility", "Pass"),
              ], 9.2)
    add_bullets(slide, [
        "Testing covered CRUD behavior, authentication, order creation, status updates, and tracking visibility.",
        "Screenshots and evidence pages demonstrate expected outcomes for unit and system test cases.",
    ], Inches(0.85), Inches(5.25), Inches(11.6), Inches(1.0), 15)

    # 14. Outcome
    slide = make_slide(prs, "Outcome and Lessons Learned")
    add_bullets(slide, [
        "A modular PHP/MySQL food delivery platform was implemented with customer and admin workflows.",
        "Database normalization, prepared statements, validation, password hashing, sessions, and CSRF protection strengthened reliability and security.",
        "Responsive design improved usability across devices.",
        "Configuration through environment variables improved maintainability across local and deployment environments.",
    ], Inches(0.95), Inches(1.35), Inches(11.4), Inches(4.6), 18)

    # 15. Future Recommendations
    slide = make_slide(prs, "Future Recommendations")
    future = [
        ("Mobile Apps", "Native Android/iOS apps with device GPS and push notifications."),
        ("Analytics", "Demand forecasting, sales dashboards, and delivery performance metrics."),
        ("Real-Time Delivery", "Live rider GPS integration instead of simulated route progress."),
        ("Marketplace Growth", "Extend from a single-restaurant model to multi-restaurant onboarding."),
    ]
    for idx, (title, body) in enumerate(future):
        x = Inches(0.85 + (idx % 2) * 6.1)
        y = Inches(1.45 + (idx // 2) * 2.05)
        add_card(slide, x, y, Inches(5.45), Inches(1.45), title, body, WHITE, [TEAL, ORANGE, GREEN, BLUE][idx])

    # 16. Closing
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    thanks = slide.shapes.add_textbox(Inches(1.15), Inches(2.25), Inches(11), Inches(1.0))
    set_text(thanks.text_frame, "Thank You", 46, True, WHITE, PP_ALIGN.CENTER)
    sub = slide.shapes.add_textbox(Inches(2.1), Inches(3.35), Inches(9.1), Inches(0.7))
    set_text(sub.text_frame, "FOOD SEWA - Online Food Ordering and Delivery Management System", 18, False, MINT, PP_ALIGN.CENTER)
    add_footer(slide, len(prs.slides))

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
